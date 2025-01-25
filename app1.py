import pandas as pd
from docx import Document
datafile=pd.ExcelFile("Book.xlsx")
subject=["MA3354","CS3301","CS3351","CS3352","CS3391"]
doc = Document("finaldoc.docx")
j=0
num=0
semester=4
samp=[]
def table1(sub,n):
    global samp
    global num
    global j
    df=pd.read_excel("Book.xlsx",sheet_name=sub)
    dic1=[]
    dic1.append(n+1)
    dic1.append(sub)
    dic1+=[sub]
    dic1+=[1]
    dic1+=[2]
    dic1+=[3]
    dic1+=[4]
    dic1+=["dfgbn"]
    dic1+=["fds"]
    dic1+=[len(df["name"])]
    dic1+=[len(df[df["UR"]!="ab"])]
    dic1+=[dic1[9]-dic1[10]]
    dic1+=[len(df[(df["UR"] != 'ab') & (df["UR"] != 'RA')])]
    dic1+=[dic1[10]-dic1[12]]
    dic1+=[f"{(dic1[12]/dic1[10])*100:.1f}"]
    if j==0:
        oac=[]
        for sheet_name in datafile.sheet_names:
            dfs = pd.read_excel("Book.xlsx", sheet_name=sheet_name) 
            if 'UR' in dfs.columns and 'name' in dfs.columns:
                filtered_names = dfs[(dfs['UR'] != 'ab') & (dfs['UR'] != 'RA')]['name'].tolist()
            oac.append(filtered_names)
        samp=oac[0]
        for sl in range(1,len(oac)):
            samp1=[item for item in samp if item in oac[sl]]
            samp=samp1
        dfs = pd.read_excel("Book.xlsx", sheet_name="arrear")
        num=len(dfs[dfs["Arrear history"]!=0])
    j=1
    dic1+=[f"{len(samp):.2f}"]
    dic1+=[int((len(samp)/len(df["name"]))*100)]#f"{(len(samp)/len(df["name"]))*100:.2f}"
    dic1+=[num]
    print(dic1)
    return dic1


def table2(sub,n):
    df=pd.read_excel("Book.xlsx",sheet_name=sub)
    dic1=[]
    dic1.append(sub)
    dic1+=[sub]
    dic1+=[1]
    dic1+=[2]
    dic1+=[3]
    dic1+=[4]
    dic1+=["dfgbn"]
    dic1+=["fds"]
    dic1+=[len(df["name"]),len(df["name"]),len(df["name"])]*3
    exm=["IA1","IA2","UR"]
    for i in range(2):
        dic1.append(len(df[(df[exm[i]]!="ab") & (df[exm[i]]!=0)]))
    for i in range(3):
        dic1.append(len(df["name"])-len(df[(df[exm[i]]!="ab") & (df[exm[i]]!=0)]))
    for i in range(2):
        dic1.append(len(df[(df[exm[i]] != 0) & (df[exm[i]]>35)]))
    dic1.append(len(df[(df["UR"] != 'ab') & (df["UR"] != 'RA')]))
    for i in range(2):
        dic1.append(len(df[(df[exm[i]] == 0) & (df[exm[i]]<35)]))
    dic1.append(len(df[(df["UR"] == 'ab') & (df["UR"] == 'RA')]))
    for i in range(2):
        dic1.append((len(df[(df[exm[i]] != 0) & (df[exm[i]]>35)])/len(df["name"]))*100)
    dic1.append((len(df[(df["UR"] != 'ab') & (df["UR"] != 'RA')])/len(df["name"]))*100)
    print(dic1)
    return dic1

def table3(table, datafile):
    # Load the "arrear" sheet
    df = pd.read_excel(datafile, sheet_name="arrear")

    # Filter students with arrears
    arrear_students = df[df["Arrear history"] == "ar"]

    # Identify available semester columns
    semester_columns = [col for col in df.columns if col.startswith("sem")]
    max_semesters = len(semester_columns)

    # Prepare rows for each student
    rows_data = []
    idx=1
    for idxx, row in arrear_students.iterrows():
        # Calculate total arrears
        total_arrears = sum([row.get(sem, 0) for sem in semester_columns])

        # Create student data dictionary
        student_data = {
            "SL.NO.": idx,
            "REGISTER NUMBER": row.get("regno", "-"),
            "NAME OF THE STUDENT": row.get("name", "-"),
            "NO. OF ARREARS": total_arrears,
            "SEMESTERS": [
                row.get(f"sem{i}", "-") for i in range(1, 9)
            ],  # Fill missing semesters with "-"
        }
        idx+=1
        rows_data.append(student_data)
        

    # Ensure table has enough rows
    while len(table.rows) < len(rows_data) + 2:  # Add rows after header
        table.add_row()

    # Fill the table with data
    for i, student_data in enumerate(rows_data):
        row = table.rows[i + 2]  # Offset for header rows
        row.cells[0].text = str(student_data["SL.NO."])
        row.cells[1].text = str(student_data["REGISTER NUMBER"])
        row.cells[2].text = str(student_data["NAME OF THE STUDENT"])
        row.cells[3].text = str(student_data["NO. OF ARREARS"])
        for sem_idx, sem_value in enumerate(student_data["SEMESTERS"]):
            if sem_idx < len(row.cells) - 4:  # Avoid index issues
                row.cells[4 + sem_idx].text = str(sem_value)

    print("Table 3 has been updated successfully.")


def table4(table, datafile):
    """
    Fill the "COURSE-WISE ARREAR DETAILS REPORT" table in the Word document.
    - `table`: The Word table object to be filled.
    - `datafile`: The Excel file containing arrear details.
    """
    # Load Excel file and get sheet names
    xl = pd.ExcelFile(datafile)
    sheet_names = xl.sheet_names
    sheet_names.pop()
    rows_data = []
    
    # Iterate over each sheet
    for sheet_name in sheet_names:
        # Load the sheet data
        df = xl.parse(sheet_name)

        # Extract semester number (last character of sheet name)
        sem_no = sheet_name[-1] if sheet_name[-1].isdigit() else "N/A"

        # Count arrears and absent
        arrear_students = df[(df["UR"] == "RA") | (df["UR"] == "ab")]
        arrear_count = len(arrear_students)

        # Prepare names and reg. numbers as a single string
        student_list = "\n".join(
            [
                f"{row['name']} ({row['regno']})"
                for _, row in arrear_students.iterrows()
            ]
        )

        # Extract course information
        course_code = sheet_name  # Sheet name as course code
        course_title = df.columns[0] if len(df.columns) > 0 else "N/A"
        L, T, P, C = (1, 2, 3, 4)  # Default values; adjust if needed

        # Prepare data for a row
        rows_data.append({
            "SL.NO.": len(rows_data) + 1,
            "COURSE CODE": course_code,
            "COURSE TITLE": course_title,
            "L":L,
            "T":T,
            "P":P,
            "C":C,
            "SEM NO.": sem_no,
            "NO. OF ARREARS": arrear_count,
            "NAME AND REGISTERNO. OF THE STUDENTS": student_list,
            "ACTION PLAN": "",
            "EXECUTION CONFIRMATION": ""
        })

    # Ensure table has enough rows
    while len(table.rows) < len(rows_data) + 1:  # Add rows after header
        table.add_row()

    # Fill the table
    for i, row_data in enumerate(rows_data):
        row = table.rows[i + 1]  # Offset for header row
        row.cells[0].text = str(row_data["SL.NO."])
        row.cells[1].text = str(row_data["COURSE CODE"])
        row.cells[2].text = str(row_data["COURSE TITLE"])
        row.cells[3].text = str(row_data["L"])
        row.cells[4].text = str(row_data["T"])
        row.cells[5].text = str(row_data["P"])
        row.cells[6].text = str(row_data["C"])
        row.cells[7].text = str(row_data["SEM NO."])
        row.cells[8].text = str(row_data["NO. OF ARREARS"])
        row.cells[9].text = str(row_data["NAME AND REGISTERNO. OF THE STUDENTS"])
        row.cells[10].text = str(row_data["ACTION PLAN"])
        row.cells[11].text = str(row_data["EXECUTION CONFIRMATION"])

    print("Table 4 has been updated successfully.")

def table5(sub,n):
     df=pd.read_excel("Book.xlsx",sheet_name=sub)
     data=[]
     data.append(n)
     data.append(sub)
     data.append(sub)
     data.append("ctype")
     data.append("ctype")
     data.append(1)
     data.append(2)
     data.append(3)
     data.append(4)
     data.append(len(df[df["UR"]=='O']))
     data.append(len(df[df["UR"]=='A+']))
     data.append(len(df[df["UR"]=='A']))
     data.append(len(df[df["UR"]=='B+']))
     data.append(len(df[df["UR"]=='B']))
     data.append(len(df[df["UR"]=='C']))
     return data


def table6(n):
     df=pd.read_excel("Book.xlsx",sheet_name="arrear")
     data=[]
     










flg=0
for table_index, table in enumerate(doc.tables): #to iterate through each table

    if len(table.rows)>2:
        if len(table.rows)-2 < len(subject) and (table_index!=4):
            print(len(table.rows),len(subject))
            for _ in range(len(subject) - (len(table.rows)-2)):
                table.add_row()#to add rows
                print("tab index",table_index)
        i=0
        for row_index, row in enumerate(table.rows):#to iterate through row
            if row_index>=2:#to ignore sign table
                if table_index==0:
                    for col_in in range(15, 18):
                        start_cell = table.cell(2, col_in) 
                        for row_in in range(2, len(table.rows)): 
                            start_cell.merge(table.cell(row_in, col_in))
                    if (i<len(subject)):
                        data=table1(subject[i],i)
                    flg=1
                elif table_index==2:
                    if (i<len(subject)):
                        data=table2(subject[i],i)
                    flg=1
                elif table_index==8:
                    if (i<len(subject)):
                        data=table5(subject[i],i)
                    flg=1
                    print(len(table.rows))
                else:
                    print("else part")
                if flg==1:
                    for cell_index, cell in enumerate(row.cells): 
                        cell.text = str(data[cell_index])
                    flg=0
                i+=1

    if table_index==4:
        table3(table,"Book.xlsx")
    elif table_index==6:
        table4(table,"Book.xlsx")
    else:
        print("last else") 
    if table_index>=9:
        break
doc.save("output_checked_cells.docx")